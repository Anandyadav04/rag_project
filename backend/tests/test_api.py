import os
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.config import settings
from app.database import Base, get_db
from app.main import app
from app.services.extractor import DocumentExtractor
from app.services.chunker import LegalChunker
from app.services.embedder import EmbeddingService

SQLALCHEMY_TEST_DATABASE_URL = "sqlite:///:memory:"
test_engine = create_engine(
    SQLALCHEMY_TEST_DATABASE_URL,
    connect_args={"check_same_thread": False},
    poolclass=StaticPool
)
TestingSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=test_engine)

def override_get_db():
    db = TestingSessionLocal()
    try:
        yield db
    finally:
        db.close()

app.dependency_overrides[get_db] = override_get_db
client = TestClient(app)

@pytest.fixture(autouse=True)
def setup_test_db():
    Base.metadata.create_all(bind=test_engine)
    yield
    Base.metadata.drop_all(bind=test_engine)

def create_sample_docx(file_path: str):
    import docx
    doc = docx.Document()
    doc.add_heading("NON-DISCLOSURE AGREEMENT", level=1)
    doc.add_paragraph("This Non-Disclosure Agreement ('Agreement') is entered into by and between Party A and Party B.")
    doc.add_heading("SECTION 1. DEFINITIONS", level=2)
    doc.add_paragraph("Confidential Information means all non-public, proprietary information disclosed by either party.")
    doc.add_heading("SECTION 2. CONFIDENTIALITY OBLIGATIONS", level=2)
    doc.add_paragraph("The receiving party agrees to hold all Confidential Information in strict confidence.")
    doc.save(file_path)

def create_sample_pdf(file_path: str):
    import fitz
    doc = fitz.open()
    page1 = doc.new_page()
    page1.insert_text((50, 50), "MASTER SERVICES AGREEMENT\n\nSECTION 1. TERM AND TERMINATION\nThis Agreement shall commence on the Effective Date and remain in effect for 3 years.")
    page2 = doc.new_page()
    page2.insert_text((50, 50), "SECTION 2. INDEMNIFICATION AND LIABILITY\nEach party shall indemnify and hold harmless the other party against any third-party claims.")
    doc.save(file_path)
    doc.close()

def test_extractor_pdf(tmp_path):
    pdf_path = str(tmp_path / "sample.pdf")
    create_sample_pdf(pdf_path)

    pages, total_pages = DocumentExtractor.extract(pdf_path, "pdf")
    assert total_pages == 2
    assert pages[0].page_number == 1
    assert "MASTER SERVICES AGREEMENT" in pages[0].text
    assert pages[1].page_number == 2
    assert "INDEMNIFICATION" in pages[1].text

def test_extractor_docx(tmp_path):
    docx_path = str(tmp_path / "sample.docx")
    create_sample_docx(docx_path)

    pages, total_pages = DocumentExtractor.extract(docx_path, "docx")
    assert total_pages >= 1
    assert "NON-DISCLOSURE AGREEMENT" in pages[0].text
    assert len(pages[0].sections_detected) > 0

def test_legal_chunker(tmp_path):
    pdf_path = str(tmp_path / "sample.pdf")
    create_sample_pdf(pdf_path)
    pages, _ = DocumentExtractor.extract(pdf_path, "pdf")

    chunker = LegalChunker(target_chunk_size=150, overlap=30)
    chunks = chunker.create_chunks(pages)
    
    assert len(chunks) > 0
    assert "chunk_index" in chunks[0]
    assert "page_number" in chunks[0]
    assert "section_heading" in chunks[0]
    assert chunks[0]["chunk_index"] == 0

def test_embedding_service():
    embedder = EmbeddingService()
    vecs = embedder.generate_embeddings(["This is a legal clause.", "Indemnification requirement."])
    assert len(vecs) == 2
    assert len(vecs[0]) == settings.EMBEDDING_DIMENSION

def test_document_api_lifecycle(tmp_path):
    docx_path = str(tmp_path / "nda_test.docx")
    create_sample_docx(docx_path)

    # 1. Upload API (POST /documents/upload)
    with open(docx_path, "rb") as f:
        response = client.post("/documents/upload", files={"file": ("nda_test.docx", f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    
    assert response.status_code == 201
    upload_res = response.json()
    doc_id = upload_res["document"]["id"]
    assert upload_res["document"]["filename"] == "nda_test.docx"

    # 2. List Documents API (GET /documents)
    response = client.get("/documents")
    assert response.status_code == 200
    docs = response.json()
    assert len(docs) >= 1
    assert any(d["id"] == doc_id for d in docs)

    # 3. Get Document Detail API (GET /documents/{id})
    response = client.get(f"/documents/{doc_id}")
    assert response.status_code == 200
    doc_detail = response.json()
    assert doc_detail["id"] == doc_id
    assert doc_detail["status"] in ["uploaded", "processing", "completed"]

    # 4. Get Document Pages API (GET /documents/{id}/pages)
    response = client.get(f"/documents/{doc_id}/pages")
    assert response.status_code == 200
    pages_res = response.json()
    assert pages_res["document_id"] == doc_id

    # 5. Delete Document API (DELETE /documents/{id})
    response = client.delete(f"/documents/{doc_id}")
    assert response.status_code == 200
    assert response.json()["message"] == f"Document '{doc_id}' deleted successfully"

    # Verify document no longer exists
    response = client.get(f"/documents/{doc_id}")
    assert response.status_code == 404
